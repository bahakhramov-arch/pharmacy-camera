# -*- coding: utf-8 -*-
"""Чтение журнала событий и построение Excel-отчётов (день / неделя / месяц)."""

import csv
import datetime as dt
import os
from collections import defaultdict

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WEEKDAYS = ["пн", "вт", "ср", "чт", "пт", "сб", "вс"]


# ---------------------------------------------------------------- периоды

def workday(when, day_start=0):
    """Рабочий день, к которому относится момент времени.
    day_start - час, с которого начинается новый рабочий день (0 = полночь).
    Для ночных смен ставят, например, 5: тогда 02:00 ещё относится к вчерашнему дню."""
    return (when - dt.timedelta(hours=day_start)).date()


def week_range(d):
    start = d - dt.timedelta(days=d.weekday())
    return start, start + dt.timedelta(days=6)


def month_range(d):
    start = d.replace(day=1)
    nxt = (start + dt.timedelta(days=32)).replace(day=1)
    return start, nxt - dt.timedelta(days=1)


# ---------------------------------------------------------------- журнал

def read_visitors(log_file):
    """Колонки: Дата, Время, Филиал, Камера, Событие."""
    visits = []
    if not os.path.exists(log_file):
        return visits
    with open(log_file, encoding="utf-8-sig") as f:
        for row in csv.reader(f):
            if len(row) < 5 or row[0] == "Дата":
                continue
            try:
                when = dt.datetime.strptime(f"{row[0]} {row[1]}", "%Y-%m-%d %H:%M:%S")
            except ValueError:
                continue
            visits.append({"when": when, "branch": row[2], "camera": row[3],
                           "type": row[4]})
    visits.sort(key=lambda v: v["when"])
    return visits


def visitor_stats(visits, start, end, day_start=0):
    """Сводка по посетителям: по дням, по часам, всего за период.
    total / by_* - входы (сколько человек зашло); exits / out_by_* - выходы."""
    by_day = defaultdict(int)
    by_hour = defaultdict(int)
    by_hour_day = defaultdict(lambda: defaultdict(int))
    out_by_day = defaultdict(int)
    out_by_hour = defaultdict(int)
    total = exits = 0
    for v in visits:
        d = workday(v["when"], day_start)
        if not start <= d <= end:
            continue
        if v["type"] == "выход":
            out_by_day[d] += 1
            out_by_hour[v["when"].hour] += 1
            exits += 1
        else:
            by_day[d] += 1
            by_hour[v["when"].hour] += 1
            by_hour_day[d][v["when"].hour] += 1
            total += 1
    days_n = max(1, (end - start).days + 1)
    peak_hour = max(by_hour, key=by_hour.get) if by_hour else None
    busiest_day = max(by_day, key=by_day.get) if by_day else None
    return {
        "total": total,
        "exits": exits,
        "avg_per_day": round(total / days_n, 1),
        "by_day": dict(by_day),
        "out_by_day": dict(out_by_day),
        "by_hour": {h: by_hour.get(h, 0) for h in range(24)},
        "out_by_hour": {h: out_by_hour.get(h, 0) for h in range(24)},
        "by_hour_day": {d: dict(hours) for d, hours in by_hour_day.items()},
        "peak_hour": peak_hour,
        "busiest_day": busiest_day,
        "max_day": by_day.get(busiest_day, 0) if busiest_day else 0,
    }


def read_events(log_file):
    """Колонки: Дата, Время, Сотрудник, Филиал, Событие, Опоздание."""
    events = []
    if not os.path.exists(log_file):
        return events
    with open(log_file, encoding="utf-8-sig") as f:
        for row in csv.reader(f):
            if len(row) < 6 or row[0] == "Дата":
                continue
            try:
                when = dt.datetime.strptime(f"{row[0]} {row[1]}", "%Y-%m-%d %H:%M:%S")
                late = int(row[5] or 0)
            except ValueError:
                continue
            events.append({"when": when, "name": row[2], "branch": row[3],
                           "type": row[4], "late": late})
    events.sort(key=lambda e: e["when"])
    return events


def build_days(events, start, end, last_seen=None, now=None, day_start=0):
    """Для каждой пары (рабочий день, сотрудник) считает приход, уход и отработанное время.
    Время считается по парам «приход -> уход», перерывы вычитаются.

    Поля результата:
      unclosed - в конце последнего прихода нет ухода;
      active   - то же, но это текущий рабочий день (человек, скорее всего, ещё на работе)."""
    groups = defaultdict(list)
    for e in events:
        d = workday(e["when"], day_start)
        if start <= d <= end:
            groups[(d, e["name"])].append(e)

    today_wd = workday(now, day_start) if now is not None else None

    days = {}
    for (d, name), evs in groups.items():
        worked = 0.0
        open_start = None
        first_in = None
        last_out = None
        late = 0
        for e in evs:
            if e["type"] == "приход":
                if open_start is None:
                    open_start = e["when"]
                if first_in is None:
                    first_in = e["when"]
                    late = e["late"]
            elif e["type"] == "уход":
                if open_start is not None:
                    worked += (e["when"] - open_start).total_seconds()
                    open_start = None
                last_out = e["when"]
        if first_in is None:
            continue
        unclosed = False
        active = False
        if open_start is not None:
            unclosed = True
            ls = (last_seen or {}).get(name)
            end_t = None
            if ls is not None and workday(ls, day_start) == d and ls > open_start:
                end_t = ls
            elif ls is None and today_wd is not None and d == today_wd:
                end_t = now
            if end_t is not None:
                worked += (end_t - open_start).total_seconds()
            active = (d == today_wd)
        days[(d, name)] = {
            "branch": evs[0]["branch"], "first_in": first_in, "last_out": last_out,
            "worked_sec": worked, "late": late, "unclosed": unclosed, "active": active,
        }
    return days


# ---------------------------------------------------------------- Excel

_HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_WEEKEND_FILL = PatternFill("solid", fgColor="FCE4D6")
_TOTAL_FILL = PatternFill("solid", fgColor="DDEBF7")
_THIN = Side(style="thin", color="BBBBBB")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _head(ws, row, values):
    for i, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.fill, c.font, c.alignment, c.border = _HEAD_FILL, _HEAD_FONT, _CENTER, _BORDER


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _staff_sheets(wb, days, start, end, title, now):
    """Листы по сотрудникам: Сводка, По дням, Табель."""
    names = sorted({n for (_, n) in days})

    # ---------- Лист 1: Сводка ----------
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Сформирован: {(now or dt.datetime.now()):%d.%m.%Y %H:%M}"
    ws["A2"].font = Font(italic=True, color="666666")

    cols = ["Сотрудник", "Филиал", "Дней на работе", "Всего (ч:мм)", "Всего часов",
            "В среднем за день (ч)", "Опозданий (раз)", "Опоздание всего (мин)",
            "Дней без отметки ухода"]
    _head(ws, 4, cols)
    ws.row_dimensions[4].height = 32

    agg = {}
    for (d, n), info in days.items():
        a = agg.setdefault(n, {"branch": info["branch"], "days": 0, "sec": 0.0,
                               "late_cnt": 0, "late_min": 0, "unclosed": 0})
        a["days"] += 1
        a["sec"] += info["worked_sec"]
        if info["late"] > 0:
            a["late_cnt"] += 1
            a["late_min"] += info["late"]
        # текущая смена (человек ещё на работе) - не считается пропущенным уходом
        if info["unclosed"] and not info["active"]:
            a["unclosed"] += 1

    row = 5
    tot = {"days": 0, "sec": 0.0, "late_cnt": 0, "late_min": 0}
    for n in names:
        a = agg[n]
        hours = a["sec"] / 3600
        values = [n, a["branch"], a["days"], a["sec"] / 86400, round(hours, 2),
                  round(hours / a["days"], 2) if a["days"] else 0,
                  a["late_cnt"], a["late_min"], a["unclosed"]]
        for i, v in enumerate(values, start=1):
            c = ws.cell(row=row, column=i, value=v)
            c.border = _BORDER
            if i == 4:
                c.number_format = "[h]:mm"
            if i >= 3:
                c.alignment = Alignment(horizontal="center")
        for k in ("days", "sec", "late_cnt", "late_min"):
            tot[k] += a[k]
        row += 1
    total_vals = ["ИТОГО", "", tot["days"], tot["sec"] / 86400, round(tot["sec"] / 3600, 2),
                  "", tot["late_cnt"], tot["late_min"], ""]
    for i, v in enumerate(total_vals, start=1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.fill, c.border = Font(bold=True), _TOTAL_FILL, _BORDER
        if i == 4:
            c.number_format = "[h]:mm"
        if i >= 3:
            c.alignment = Alignment(horizontal="center")
    _widths(ws, [28, 16, 14, 14, 12, 16, 14, 16, 16])
    ws.freeze_panes = "A5"

    # ---------- Лист 2: По дням ----------
    ws2 = wb.create_sheet("По дням")
    _head(ws2, 1, ["Дата", "День", "Сотрудник", "Филиал", "Пришёл", "Ушёл",
                   "Отработано (ч:мм)", "Опоздание (мин)", "Примечание"])
    ws2.row_dimensions[1].height = 30
    r = 2
    for (d, n) in sorted(days, key=lambda k: (k[0], k[1])):
        info = days[(d, n)]
        note = ""
        out_val = info["last_out"].time().replace(second=0, microsecond=0) if info["last_out"] else None
        if info["active"]:
            note = "Сейчас на работе (время считается по последней отметке)"
            out_val = "—"
        elif info["unclosed"]:
            note = "Уход не зафиксирован (считано по последней отметке)"
            if out_val is None:
                out_val = "—"
        values = [d, WEEKDAYS[d.weekday()], n, info["branch"],
                  info["first_in"].time().replace(second=0, microsecond=0), out_val,
                  info["worked_sec"] / 86400, info["late"], note]
        for i, v in enumerate(values, start=1):
            c = ws2.cell(row=r, column=i, value=v)
            c.border = _BORDER
            if i == 1:
                c.number_format = "DD.MM.YYYY"
            elif i in (5, 6):
                c.number_format = "HH:MM"
                c.alignment = Alignment(horizontal="center")
            elif i == 7:
                c.number_format = "[h]:mm"
                c.alignment = Alignment(horizontal="center")
            elif i in (2, 8):
                c.alignment = Alignment(horizontal="center")
            if d.weekday() >= 5:
                c.fill = _WEEKEND_FILL
        r += 1
    _widths(ws2, [13, 7, 28, 16, 10, 10, 16, 15, 52])
    ws2.freeze_panes = "A2"

    # ---------- Лист 3: Табель (часы по дням) ----------
    ws3 = wb.create_sheet("Табель")
    all_days = []
    d = start
    while d <= end:
        all_days.append(d)
        d += dt.timedelta(days=1)
    ws3.cell(row=1, column=1, value="Сотрудник")
    ws3.cell(row=2, column=1, value="")
    for j, d in enumerate(all_days, start=2):
        ws3.cell(row=1, column=j, value=d.day)
        ws3.cell(row=2, column=j, value=WEEKDAYS[d.weekday()])
    total_col = len(all_days) + 2
    ws3.cell(row=1, column=total_col, value="Итого, ч")
    ws3.cell(row=2, column=total_col, value="")
    for rr in (1, 2):
        for cc in range(1, total_col + 1):
            c = ws3.cell(row=rr, column=cc)
            c.fill, c.font, c.alignment, c.border = _HEAD_FILL, _HEAD_FONT, _CENTER, _BORDER
    for j, d in enumerate(all_days, start=2):
        if d.weekday() >= 5:
            for rr in (1, 2):
                ws3.cell(row=rr, column=j).fill = PatternFill("solid", fgColor="C55A11")

    rr = 3
    for n in names:
        c = ws3.cell(row=rr, column=1, value=n)
        c.border = _BORDER
        total = 0.0
        for j, d in enumerate(all_days, start=2):
            info = days.get((d, n))
            cell = ws3.cell(row=rr, column=j)
            cell.border = _BORDER
            cell.alignment = Alignment(horizontal="center")
            if d.weekday() >= 5:
                cell.fill = _WEEKEND_FILL
            if info:
                h = round(info["worked_sec"] / 3600, 2)
                cell.value = h
                total += h
        tc = ws3.cell(row=rr, column=total_col, value=round(total, 2))
        tc.font, tc.fill, tc.border = Font(bold=True), _TOTAL_FILL, _BORDER
        tc.alignment = Alignment(horizontal="center")
        rr += 1
    ws3.column_dimensions["A"].width = 28
    for j in range(2, total_col):
        ws3.column_dimensions[get_column_letter(j)].width = 6
    ws3.column_dimensions[get_column_letter(total_col)].width = 11
    ws3.freeze_panes = "B3"



def _clients_sheet(ws, st, start, end, title, now):
    """Лист «Клиенты»: итог за период, по дням и по часам (вход / выход)."""
    ws["A1"] = f"Клиенты. {title}"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Сформирован: {(now or dt.datetime.now()):%d.%m.%Y %H:%M}. "
                f"Подсчёт по линии у двери; «внутри» = вход − выход (оценка).")
    ws["A2"].font = Font(italic=True, color="666666")

    peak = st["peak_hour"]
    busiest = st["busiest_day"]
    _head(ws, 4, ["Зашло", "Вышло", "В среднем за день (вход)", "Час пик",
                  "Самый загруженный день"])
    ws.row_dimensions[4].height = 32
    vals = [st["total"], st["exits"], st["avg_per_day"],
            f"{peak:02d}:00 ({st['by_hour'][peak]} чел.)" if peak is not None else "—",
            f"{busiest:%d.%m.%Y} ({st['max_day']} чел.)" if busiest else "—"]
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=5, column=i, value=v)
        c.border, c.alignment, c.font = _BORDER, _CENTER, Font(bold=True)

    # --- по дням
    _head(ws, 7, ["Дата", "День", "Зашло", "Вышло", "Разница"])
    r = 8
    d = start
    while d <= end:
        n_in, n_out = st["by_day"].get(d, 0), st["out_by_day"].get(d, 0)
        for i, v in enumerate([d, WEEKDAYS[d.weekday()], n_in, n_out, n_in - n_out], start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.border = _BORDER
            if i == 1:
                c.number_format = "DD.MM.YYYY"
            else:
                c.alignment = Alignment(horizontal="center")
            if d.weekday() >= 5:
                c.fill = _WEEKEND_FILL
        r += 1
        d += dt.timedelta(days=1)
    for i, v in enumerate(["ИТОГО", "", st["total"], st["exits"], st["total"] - st["exits"]],
                          start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.font, c.fill, c.border = Font(bold=True), _TOTAL_FILL, _BORDER
        if i > 1:
            c.alignment = Alignment(horizontal="center")

    # --- по часам (справа от таблицы по дням)
    for i, v in enumerate(["Час", "Зашло", "Вышло"], start=7):
        c = ws.cell(row=7, column=i, value=v)
        c.fill, c.font, c.alignment, c.border = _HEAD_FILL, _HEAD_FONT, _CENTER, _BORDER
    hours = [h for h in range(24) if st["by_hour"][h] or st["out_by_hour"][h]]
    for k, h in enumerate(hours):
        for i, v in enumerate([f"{h:02d}:00–{h:02d}:59", st["by_hour"][h],
                               st["out_by_hour"][h]], start=7):
            c = ws.cell(row=8 + k, column=i, value=v)
            c.border, c.alignment = _BORDER, Alignment(horizontal="center")
            if h == peak:
                c.fill = _TOTAL_FILL
    if hours:
        chart = BarChart()
        chart.title = "Клиенты по часам"
        chart.y_axis.title = "человек"
        chart.add_data(Reference(ws, min_col=8, max_col=9, min_row=7, max_row=7 + len(hours)),
                       titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=7, min_row=8, max_row=7 + len(hours)))
        chart.height, chart.width = 8, 18
        ws.add_chart(chart, "K7")

    _widths(ws, [14, 12, 16, 16, 24, 3, 14, 10, 10])
    ws.freeze_panes = "A6"


def make_report(path, events, start, end, title, last_seen=None, now=None, day_start=0,
                visits=None):
    """Создаёт Excel: листы по сотрудникам и лист «Клиенты» (вход/выход).
    Возвращает False, если за период нет ни отметок сотрудников, ни посетителей."""
    days = build_days(events, start, end, last_seen, now, day_start)
    vstats = visitor_stats(visits or [], start, end, day_start)
    has_visits = bool(vstats["total"] or vstats["exits"])
    if not days and not has_visits:
        return False
    wb = Workbook()
    if days:
        _staff_sheets(wb, days, start, end, title, now)
    if has_visits:
        ws = wb.create_sheet("Клиенты") if days else wb.active
        if not days:
            ws.title = "Клиенты"
        _clients_sheet(ws, vstats, start, end, title, now)
    wb.save(path)
    return True
